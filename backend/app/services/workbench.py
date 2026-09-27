"""Работа с подключённой базой: SQL-консоль, вопросы на русском (text-to-SQL), аудит структуры,
медленные запросы и описание базы (словарь данных в Word).

Всё выполняется так же, как анализ запросов: только одиночные SELECT (Safety Engine), в READ ONLY
транзакции с таймаутом, пользователем с правом на чтение.
"""

from __future__ import annotations

import base64
import datetime as dt
import decimal
import io
import json
import math
import re
import time
import uuid

from app.config import get_settings
from app.models import Dialect, SchemaInfo, TableInfo
from app.services import safety
from app.services.ai.optimizer import extract_json
from app.services.ai.providers import LLMError, get_registry
from app.services.connectors import Connector, DBError
from app.services.sql_parser import SQLParseError, parse_statements

CONSOLE_ROW_LIMIT = 500


class WorkbenchError(Exception):
    pass


# ---------------------------------------------------------------- консоль
def _json_value(v):
    """Значение из драйвера СУБД → то, что можно отдать в JSON и показать в таблице."""
    if v is None or isinstance(v, (bool, int, str)):
        return v
    if isinstance(v, float):
        return v if math.isfinite(v) else str(v)
    if isinstance(v, decimal.Decimal):
        return float(v) if v.is_finite() and abs(v) < 2 ** 53 else str(v)
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, dt.timedelta):
        return str(v)
    if isinstance(v, (bytes, bytearray, memoryview)):
        b = bytes(v)
        return "0x" + b[:32].hex() + ("…" if len(b) > 32 else "")
    if isinstance(v, uuid.UUID):
        return str(v)
    if isinstance(v, (list, dict)):
        return json.dumps(v, ensure_ascii=False, default=str)
    return str(v)


def _with_limit(sql: str, dialect: Dialect, limit: int) -> str:
    """Добавляет LIMIT, если в запросе верхнего уровня его нет: консоль показывает первые строки,
    и тянуть из базы миллионы строк ради них незачем."""
    sql = sql.strip().rstrip(";").strip()
    try:
        tree = parse_statements(sql, dialect)[0]
    except (SQLParseError, IndexError):
        return sql
    if tree.args.get("limit") is None and tree.args.get("fetch") is None:
        return f"{sql}\nLIMIT {limit + 1}"
    return sql


def run_console(conn: Connector, dialect: Dialect, sql: str, limit: int = CONSOLE_ROW_LIMIT) -> dict:
    verdict = safety.check(sql, dialect)
    if not verdict.allowed:
        raise WorkbenchError("; ".join(verdict.reasons))
    executed = _with_limit(sql, dialect, limit)
    rows: list[list] = []

    def keep(row):
        if len(rows) <= limit:
            rows.append([_json_value(v) for v in row])

    res = conn.run(executed, row_consumer=keep)
    truncated = len(rows) > limit or res.rows_returned > limit
    return {"columns": res.columns, "rows": rows[:limit], "row_count": min(res.rows_returned, limit),
            "truncated": truncated, "elapsed_ms": round(res.elapsed_ms, 2), "executed_sql": executed,
            "warnings": verdict.warnings}


# ---------------------------------------------------------------- компактная схема для модели
def _schema_for_llm(schema: SchemaInfo, max_tables: int = 60) -> list[dict]:
    out = []
    for t in schema.tables[:max_tables]:
        out.append({"table": t.name, "rows": t.row_count,
                    "columns": [f"{c.name} {c.type}{'' if c.nullable else ' NOT NULL'}" for c in t.columns],
                    "primary_key": next((i.columns for i in t.indexes if i.primary), None),
                    "foreign_keys": [f"({', '.join(f.columns)}) → {f.ref_table}({', '.join(f.ref_columns)})"
                                     for f in t.foreign_keys]})
    return out


def _complete(system: str, user: str, model_id: str | None, json_mode: bool = True):
    reg = get_registry()
    try:
        provider, model = reg.resolve(model_id)
        return provider.complete(model, system, user, get_settings().llm_temperature, json_mode=json_mode), \
            f"{provider.name}:{model}"
    except LLMError as e:
        raise WorkbenchError(str(e)) from e


# ---------------------------------------------------------------- «Спроси базу»: вопрос на русском → SQL
ASK_SYSTEM = """Ты — помощник аналитика данных. По вопросу пользователя на русском языке составь ОДИН запрос
SELECT к базе данных {dialect_name} {version}. Пользуйся только таблицами и колонками из схемы ниже, имена пиши
точно как в схеме. Запрос только читает данные (никаких INSERT/UPDATE/DELETE/DDL). Если вопрос про «топ» или
список — ограничь число строк (LIMIT). Колонкам результата давай понятные псевдонимы.

Ответь строго JSON-объектом:
{{"sql": "запрос", "explanation": "одно-два предложения по-русски: что считает запрос",
  "chart": {{"type": "bar" | "line" | "none", "x": "колонка для подписей", "y": "числовая колонка"}}}}
Диаграмму предлагай, только если результат — ряд подписей и чисел (bar) или динамика по времени (line)."""


def ask(conn: Connector, dialect: Dialect, version: str | None, schema: SchemaInfo, question: str,
        model_id: str | None) -> dict:
    if not question.strip():
        raise WorkbenchError("Задайте вопрос")
    system = ASK_SYSTEM.format(dialect_name="MySQL" if dialect == "mysql" else "PostgreSQL", version=version or "")
    user = json.dumps({"schema": _schema_for_llm(schema), "question": question}, ensure_ascii=False)
    attempts = []
    t0 = time.perf_counter()
    for attempt in range(2):  # вторая попытка — с текстом ошибки СУБД (самоисправление)
        llm, used_model = _complete(system, user, model_id)
        try:
            data = extract_json(llm.text)
            sql = str(data.get("sql", "")).strip()
        except (ValueError, AttributeError):
            data, sql = {}, ""
        if not sql:
            attempts.append({"sql": None, "error": "Модель не вернула запрос"})
            break
        try:
            result = run_console(conn, dialect, sql)
            chart = data.get("chart") if isinstance(data.get("chart"), dict) else None
            if chart and (chart.get("type") not in ("bar", "line") or chart.get("x") not in result["columns"]
                          or chart.get("y") not in result["columns"]):
                chart = None
            return {"sql": sql, "explanation": data.get("explanation", ""), "chart": chart, "result": result,
                    "model": used_model, "latency_ms": round((time.perf_counter() - t0) * 1000),
                    "attempts": attempts + [{"sql": sql, "error": None}]}
        except (WorkbenchError, DBError) as e:
            attempts.append({"sql": sql, "error": str(e)})
            user = json.dumps({"schema": _schema_for_llm(schema), "question": question, "previous_sql": sql,
                               "error": str(e), "task": "Исправь запрос с учётом ошибки"}, ensure_ascii=False)
    raise WorkbenchError("Не удалось составить рабочий запрос: " + "; ".join(a["error"] for a in attempts if a["error"]))


# ---------------------------------------------------------------- аудит структуры базы
def _ident(name: str, dialect: Dialect) -> str:
    q = "`" if dialect == "mysql" else '"'
    return ".".join(q + part.replace(q, q * 2) + q for part in name.split("."))


_DATE_NAME = re.compile(r"(^|_)(date|dt|time|created|updated|deleted|birth|дата)(_|$)|_(at|on)$", re.I)
_MONEY_NAME = re.compile(r"(price|amount|cost|sum|total|salary|balance|fee|oklad|summa|стоимость|цена|сумма)", re.I)
_TEXT_TYPE = re.compile(r"^(VARCHAR|CHARACTER VARYING|CHAR|CHARACTER|TEXT|TINYTEXT|MEDIUMTEXT)", re.I)
_FLOAT_TYPE = re.compile(r"^(FLOAT|DOUBLE|REAL)", re.I)

SEVERITY_PENALTY = {"high": 12, "medium": 5, "low": 1}


def _find(code, severity, table, title, detail, fix=None):
    return {"code": code, "severity": severity, "table": table, "title": title, "detail": detail, "fix_sql": fix}


def audit(conn: Connector, dialect: Dialect, schema: SchemaInfo) -> dict:
    findings: list[dict] = []
    for t in schema.tables:
        tn = _ident(t.name, dialect)
        short = t.name.split(".")[-1]
        # 1. первичный ключ
        if not any(i.primary for i in t.indexes):
            idcol = next((c.name for c in t.columns if c.name.lower() == "id"), None)
            findings.append(_find(
                "NO_PRIMARY_KEY", "high", t.name, "Нет первичного ключа",
                "Без первичного ключа строки нельзя однозначно адресовать, InnoDB создаёт скрытый ключ, а репликация "
                "и поиск по строке работают медленнее.",
                f"ALTER TABLE {tn} ADD PRIMARY KEY ({_ident(idcol, dialect)});" if idcol
                else f"-- добавьте в {t.name} столбец-идентификатор и объявите его первичным ключом"))
        # 2. внешние ключи без индекса
        for fk in t.foreign_keys:
            n = len(fk.columns)
            if not any([c.lower() for c in i.columns[:n]] == [c.lower() for c in fk.columns] for i in t.indexes):
                cols = ", ".join(_ident(c, dialect) for c in fk.columns)
                findings.append(_find(
                    "FK_WITHOUT_INDEX", "high", t.name, f"Внешний ключ ({', '.join(fk.columns)}) без индекса",
                    f"Соединения {short} с {fk.ref_table} и удаление строк из {fk.ref_table} будут читать "
                    f"{short} целиком.",
                    f"CREATE INDEX {_ident('idx_' + short + '_' + '_'.join(fk.columns), dialect)} ON {tn} ({cols});"))
        # 3. дублирующиеся и избыточные индексы (один — префикс другого)
        idx = [i for i in t.indexes if not i.primary]
        for a in idx:
            for b in t.indexes:
                if a is b or a.unique:
                    continue
                ca, cb = [c.lower() for c in a.columns], [c.lower() for c in b.columns]
                if ca == cb[:len(ca)] and (len(ca) < len(cb) or b.primary or b.unique or b.name < a.name):
                    drop = (f"DROP INDEX {_ident(a.name, dialect)} ON {tn};" if dialect == "mysql"
                            else f"DROP INDEX {_ident('.'.join(t.name.split('.')[:-1] + [a.name]), dialect)};")
                    findings.append(_find(
                        "REDUNDANT_INDEX", "medium", t.name, f"Лишний индекс {a.name}",
                        f"Его столбцы ({', '.join(a.columns)}) — начало индекса {b.name} ({', '.join(b.columns)}). "
                        "Он только замедляет вставки и занимает место.", drop))
                    break
        if len(t.indexes) > 8:
            findings.append(_find("TOO_MANY_INDEXES", "low", t.name, f"{len(t.indexes)} индексов на одной таблице",
                                  "Каждый индекс замедляет INSERT и UPDATE. Проверьте, все ли они нужны."))
        # 4. типы данных
        for c in t.columns:
            if _DATE_NAME.search(c.name) and _TEXT_TYPE.match(c.type):
                findings.append(_find(
                    "DATE_AS_TEXT", "medium", t.name, f"Дата в текстовом столбце {c.name} ({c.type})",
                    "Сравнения и сортировка по строке работают неверно для разных форматов, индекс по диапазону дат "
                    "не используется, а функции дат требуют преобразования.",
                    f"-- ALTER TABLE {tn} ... {_ident(c.name, dialect)} → DATE / TIMESTAMP (после проверки формата данных)"))
            if _MONEY_NAME.search(c.name) and _FLOAT_TYPE.match(c.type):
                findings.append(_find(
                    "MONEY_AS_FLOAT", "medium", t.name, f"Денежная сумма {c.name} в типе {c.type}",
                    "Числа с плавающей точкой хранят суммы приближённо: 0,1 + 0,2 ≠ 0,3. Для денег нужен DECIMAL/NUMERIC.",
                    (f"ALTER TABLE {tn} MODIFY {_ident(c.name, dialect)} DECIMAL(12,2);" if dialect == "mysql"
                     else f"ALTER TABLE {tn} ALTER COLUMN {_ident(c.name, dialect)} TYPE NUMERIC(12,2);")))
        if t.columns and all(c.nullable for c in t.columns if not any(c.name in i.columns for i in t.indexes if i.primary)) \
                and len(t.columns) >= 4:
            findings.append(_find("ALL_NULLABLE", "low", t.name, "Все столбцы допускают NULL",
                                  "Обязательные поля лучше объявить NOT NULL: это защищает данные и помогает планировщику."))

    # 5. статистика СУБД: неиспользуемые индексы и таблицы без статистики
    usage = conn.index_usage()
    by_name = {t.name: t for t in schema.tables}
    for table, index in usage or []:
        t = by_name.get(table)
        if t is None or (t.row_count or 0) < 1000:
            continue
        drop = (f"DROP INDEX {_ident(index, dialect)} ON {_ident(table, dialect)};" if dialect == "mysql"
                else f"DROP INDEX {_ident('.'.join(table.split('.')[:-1] + [index]), dialect)};")
        findings.append(_find("UNUSED_INDEX", "low", table, f"Индекс {index} ни разу не использовался",
                              "С момента запуска сервера или сброса статистики к индексу не было ни одного обращения. "
                              "Прежде чем удалять, убедитесь, что статистика собрана за типичный период работы.",
                              f"-- {drop}"))
    for table in conn.never_analyzed():
        findings.append(_find("NO_STATISTICS", "medium", table, "У планировщика нет статистики",
                              "По таблице ни разу не выполнялся ANALYZE: оценки числа строк неверны, и планы запросов "
                              "могут быть плохими.", f"ANALYZE {_ident(table, dialect)};"))

    order = {"high": 0, "medium": 1, "low": 2}
    findings.sort(key=lambda f: (order[f["severity"]], f["table"] or "", f["code"]))
    score = max(0, 100 - sum(SEVERITY_PENALTY[f["severity"]] for f in findings))
    return {"score": score, "tables": len(schema.tables), "findings": findings,
            "counts": {s: sum(1 for f in findings if f["severity"] == s) for s in ("high", "medium", "low")},
            "usage_stats": usage is not None}


# ---------------------------------------------------------------- описание базы (словарь данных)
DESCRIBE_SYSTEM = """Ты — технический писатель. По структуре таблиц базы данных кратко опиши по-русски назначение
каждой таблицы (одно предложение) и каждого столбца (несколько слов). Не выдумывай того, чего не видно из имён и связей;
если назначение неочевидно, так и напиши. Ответ — строго JSON:
{"tables": {"имя_таблицы": {"description": "...", "columns": {"столбец": "..."}}}}"""

_desc_cache: dict[tuple[int, str], dict] = {}


def describe(schema: SchemaInfo, model_id: str | None, cache_key: int) -> dict:
    """Описания таблиц и столбцов от модели; по 12 таблиц за запрос, результат кэшируется."""
    key = (cache_key, ",".join(sorted(t.name for t in schema.tables)))
    if key in _desc_cache:
        return _desc_cache[key]
    out: dict = {}
    for i in range(0, len(schema.tables), 12):
        part = SchemaInfo(tables=schema.tables[i:i + 12], source=schema.source)
        llm, _ = _complete(DESCRIBE_SYSTEM, json.dumps(_schema_for_llm(part, 12), ensure_ascii=False), model_id)
        try:
            out.update(extract_json(llm.text).get("tables", {}))
        except (ValueError, AttributeError):
            continue
    _desc_cache[key] = out
    return out


def _key_mark(t: TableInfo, col: str) -> str:
    marks = []
    if any(i.primary and col in i.columns for i in t.indexes):
        marks.append("PK")
    for fk in t.foreign_keys:
        if col in fk.columns:
            marks.append(f"FK → {fk.ref_table}")
    return ", ".join(marks)


def data_dictionary_docx(name: str, dialect: Dialect, version: str | None, schema: SchemaInfo,
                         descriptions: dict | None, er_png: str | None = None) -> bytes:
    """Словарь данных в Word по ГОСТ 2.105 / 7.32: Times New Roman, поля 30/10/20/20 мм, интервал 1,5,
    таблицы с подписью «Таблица N – …» над таблицей слева."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Mm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Mm(30), Mm(10), Mm(20), Mm(20)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Times New Roman", Pt(14)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(0)
    descriptions = descriptions or {}

    def para(text, center=False, indent=True, size=None, before=0, after=0, keep=False):
        p = doc.add_paragraph()
        r = p.add_run(text)
        if size:
            r.font.size = Pt(size)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25) if indent and not center else Cm(0)
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        p.paragraph_format.keep_with_next = keep
        return p

    def table(headers, rows, widths):
        tb = doc.add_table(rows=1, cols=len(headers))
        tb.style = "Table Grid"
        tb.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row_cells, values in ((tb.rows[0].cells, headers), *((tb.add_row().cells, r) for r in rows)):
            for cell, v, w in zip(row_cells, values, widths):
                cell.width = Cm(w)
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.first_line_indent = Cm(0)
                r = p.add_run(str(v) if v is not None else "")
                r.font.size = Pt(12)
        # шапка повторяется на каждой странице
        tr = tb.rows[0]._tr
        trpr = tr.get_or_add_trPr()
        el = trpr.makeelement(qn("w:tblHeader"), {qn("w:val"): "true"})
        trpr.append(el)
        return tb

    dbms = "MySQL" if dialect == "mysql" else "PostgreSQL"
    para(f"ОПИСАНИЕ СТРУКТУРЫ БАЗЫ ДАННЫХ «{name}»", center=True, after=12)
    para(f"Документ содержит описание таблиц базы данных «{name}» (СУБД {dbms} {version or ''}): назначение таблиц, "
         f"состав и типы полей, ключи и связи между таблицами. Всего таблиц: {len(schema.tables)}. Описание "
         f"сформировано программой AI Database Optimizer {dt.date.today().strftime('%d.%m.%Y')} по структуре "
         "подключённой базы" + (", назначение таблиц и полей предложено языковой моделью и требует проверки."
                                 if descriptions else "."))
    n = 1
    para("1 Перечень таблиц", indent=True, before=24, after=12, keep=True)
    para(f"Таблица {n} – Перечень таблиц базы данных", indent=False, keep=True)
    table(["№", "Таблица", "Назначение", "Строк"],
          [[i + 1, t.name, descriptions.get(t.name, {}).get("description", ""), t.row_count if t.row_count is not None else "—"]
           for i, t in enumerate(schema.tables)], [1.0, 4.5, 9.0, 2.5])
    n += 1
    para("2 Структура таблиц", indent=True, before=24, after=12, keep=True)
    for t in schema.tables:
        d = descriptions.get(t.name, {})
        if d.get("description"):
            para(f"Таблица {t.name}: {d['description'][0].lower() + d['description'][1:]}", before=6)
        para(f"Таблица {n} – Структура таблицы {t.name}", indent=False, keep=True, before=6)
        cols = d.get("columns", {}) if isinstance(d.get("columns"), dict) else {}
        table(["Поле", "Тип", "NULL", "Ключ", "Описание"],
              [[c.name, c.type.lower(), "да" if c.nullable else "нет", _key_mark(t, c.name), cols.get(c.name, "")]
               for c in t.columns], [3.8, 3.4, 1.4, 3.4, 5.0])
        idx = [i for i in t.indexes if not i.primary]
        if idx:
            para("Индексы: " + "; ".join(f"{i.name} ({', '.join(i.columns)}){' — уникальный' if i.unique else ''}"
                                          for i in idx) + ".", before=6)
        n += 1
    links = [(t.name, fk) for t in schema.tables for fk in t.foreign_keys]
    para("3 Связи между таблицами", indent=True, before=24, after=12, keep=True)
    if er_png:
        doc.add_picture(io.BytesIO(base64.b64decode(er_png.split(",", 1)[-1])), width=Cm(17))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        para("Рисунок 1 – Схема связей таблиц базы данных", center=True, after=12)
    if links:
        for tname, fk in links:
            para(f"– {tname} ({', '.join(fk.columns)}) ссылается на {fk.ref_table} ({', '.join(fk.ref_columns)});")
    else:
        para("Внешние ключи в базе не объявлены.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
