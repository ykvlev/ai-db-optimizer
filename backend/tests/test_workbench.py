"""Консоль, аудит структуры и словарь данных — без подключения к СУБД (схема из DDL, заглушка коннектора)."""

import datetime as dt
import decimal
import io

import pytest

from app.services import audit, workbench
from app.services.connectors import ExecResult
from app.services.schema_ddl import parse_ddl

BAD_DDL = """
CREATE TABLE dept (id INT PRIMARY KEY, name VARCHAR(100));
CREATE TABLE emp (
  id INT,
  dept_id INT,
  hire_date VARCHAR(20),
  salary FLOAT,
  name VARCHAR(100),
  FOREIGN KEY (dept_id) REFERENCES dept(id)
);
CREATE TABLE pays (id INT PRIMARY KEY, emp_id INT, amount DECIMAL(10,2), created_at DATETIME);
CREATE INDEX i_emp ON pays (emp_id);
CREATE INDEX i_emp_date ON pays (emp_id, created_at);
CREATE TABLE users (id INT PRIMARY KEY, name VARCHAR(100));
CREATE TABLE orders (id INT PRIMARY KEY, user_id INT, total DECIMAL(10,2));
"""


class FakeConn:
    """Коннектор-заглушка: возвращает заранее заданные строки и статистику."""

    def __init__(self, rows=(), unused=None):
        self.rows, self.unused, self.last_sql = list(rows), unused, None

    def run(self, sql, row_consumer=None, measure_rows_examined=False):
        self.last_sql = sql
        for r in self.rows:
            row_consumer(r)
        return ExecResult(columns=["a", "b"], elapsed_ms=1.5, rows_returned=len(self.rows))

    def index_usage(self):
        return self.unused

    def _query(self, sql, args=None):
        """Ответы на проверки аудита по данным: значения уникальны, даты в ISO, «висячих» ссылок нет."""
        if "count(DISTINCT" in sql:
            return [(3, 3, 3)]
        if "REGEXP" in sql or " ~ " in sql:
            return [(2, 2, 0, 0)]
        if "max(abs(" in sql:
            return [(1000.0,)]
        if "NOT EXISTS" in sql:
            return [(0,)]
        return []

    def never_analyzed(self):
        return []


def test_console_adds_limit_and_serializes_values():
    conn = FakeConn([(decimal.Decimal("1.50"), dt.date(2025, 3, 15)), (b"\x00\x01", None)])
    r = workbench.run_console(conn, "mysql", "SELECT a, b FROM t;", limit=10)
    assert conn.last_sql.endswith("LIMIT 11")
    assert r["rows"] == [[1.5, "2025-03-15"], ["0x0001", None]]
    assert not r["truncated"]


def test_console_keeps_user_limit_and_truncates():
    conn = FakeConn([(i, i) for i in range(5)])
    r = workbench.run_console(conn, "postgres", "SELECT a, b FROM t LIMIT 100", limit=3)
    assert "LIMIT 100" in conn.last_sql and conn.last_sql.count("LIMIT") == 1
    assert r["truncated"] and len(r["rows"]) == 3


def test_console_blocks_writes():
    with pytest.raises(workbench.WorkbenchError):
        workbench.run_console(FakeConn(), "mysql", "DELETE FROM t")


def test_audit_finds_and_verifies_problems():
    schema, _ = parse_ddl(BAD_DDL, "mysql")
    r = audit.run(FakeConn(unused=[]), "mysql", schema)
    by = {(f["code"], f["table"]): f for f in r["findings"]}
    pk = by[("NO_PRIMARY_KEY", "emp")]
    assert pk["fix_kind"] == "safe" and pk["fix_sql"] == "ALTER TABLE `emp` ADD PRIMARY KEY (`id`);"
    assert ("FK_WITHOUT_INDEX", "emp") not in by  # MySQL сам создаёт индекс под внешний ключ
    date = by[("DATE_AS_TEXT", "emp")]
    assert date["fix_sql"] == "ALTER TABLE `emp` MODIFY `hire_date` DATE;" and date["fix_kind"] == "safe"
    assert by[("MONEY_AS_FLOAT", "emp")]["fix_sql"].endswith("DECIMAL(14,2);")
    assert ("MONEY_AS_FLOAT", "pays") not in by  # DECIMAL — правильно
    assert by[("REDUNDANT_INDEX", "pays")]["fix_sql"] == "DROP INDEX `i_emp` ON `pays`;"
    fk = by[("MISSING_FOREIGN_KEY", "orders")]  # user_id без объявленной связи с users
    assert "REFERENCES `users` (`id`)" in fk["fix_sql"] and fk["fix_kind"] == "safe"
    assert all(f["why"] and f["effect"] for f in r["findings"])
    assert 0 < r["score"] < 100 and set(r["categories"]) == {"integrity", "indexes", "types", "stats"}


def test_audit_score_does_not_collapse_on_small_tables():
    # 30 маленьких таблиц без статистики и с внешними ключами без индекса — это не «ноль из ста»
    ddl = "CREATE TABLE ref (id INT PRIMARY KEY);\n" + "\n".join(
        f"CREATE TABLE t{i} (id INT PRIMARY KEY, ref_id INT, FOREIGN KEY (ref_id) REFERENCES ref(id));" for i in range(30))
    schema, _ = parse_ddl(ddl, "postgres")
    r = audit.run(FakeConn(unused=None), "postgres", schema)
    assert r["categories"]["indexes"]["issues"] == 30
    assert r["score"] >= 80  # мелкие замечания (low) снимают по 15 % с таблицы в своём направлении


def test_audit_fk_without_index_postgres():
    schema, _ = parse_ddl(BAD_DDL.replace("DATETIME", "TIMESTAMP"), "postgres")
    r = audit.run(FakeConn(unused=None), "postgres", schema)
    fix = next(f["fix_sql"] for f in r["findings"] if f["code"] == "FK_WITHOUT_INDEX")
    assert fix == 'CREATE INDEX "idx_emp_dept_id" ON "emp" ("dept_id");'


def test_data_dictionary_docx():
    from docx import Document
    schema, _ = parse_ddl(BAD_DDL, "mysql")
    data = workbench.data_dictionary_docx("hr", "mysql", "8.4", schema,
                                          {"emp": {"description": "Сотрудники", "columns": {"salary": "Оклад"}}})
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "– Структура таблицы dept" in text and "1 Общие сведения" in text
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "Оклад" in cells and "FK → dept" in cells
