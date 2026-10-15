"""Измерение пересечения научных документов.

    python overlap.py

Сравниваются содержательные части документов (до списка литературы), потому что
списки литературы пересекаются всегда и дублированием текста не являются.

Две меры:
  1. дословное совпадение предложений (от 6 слов, после нормализации);
  2. доля 8-словных скользящих фрагментов первого документа, найденных во втором.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document

SCIENCE = Path(r"C:\Users\Артём\Documents\Science")
DOCS = {
    "МГУ СИТИТО": SCIENCE / "Научные материалы/01 Статьи/МГУ — СИТИТО 2026/Яковлев — сравнение моделей SQL — СИТИТО.docx",
    "ИСП РАН": SCIENCE / "Научные материалы/01 Статьи/ИСП РАН 2026/Статья ИСП РАН — с автором.docx",
    "НИР Наука 3.5": SCIENCE / "Научные материалы/05 Заявки и конкурсы/Наука 3.5 — ВоГУ 2026/Редакция с приложением — 16 октября 2026/НИР — с приложением о сравнении моделей — Наука 3.5.docx",
}

STOP = [
    r"^СПИСОК ЛИТЕРАТУРЫ",
    r"^Список использованных источников",
    r"^Литература$",
    r"^References$",
]


def document_text(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def body_only(text: str) -> str:
    """Текст до списка литературы."""
    out = []
    for line in text.split("\n"):
        if any(re.match(p, line.strip(), re.I) for p in STOP):
            break
        out.append(line)
    return "\n".join(out)


def normalize(text: str) -> str:
    text = text.lower().replace("ё", "е")
    text = re.sub(r"[^а-яa-z0-9\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def sentences(text: str, min_words: int = 6):
    text = re.sub(r"\s+", " ", text)
    out = []
    for s in re.split(r"(?<=[.!?])\s+", text):
        n = normalize(s)
        if len(n.split()) >= min_words:
            out.append((n, s.strip()))
    return out


def shingles(text: str, width: int = 8) -> set[tuple[str, ...]]:
    words = normalize(text).split()
    return {tuple(words[i:i + width]) for i in range(max(0, len(words) - width + 1))}


def compare(name_a: str, text_a: str, name_b: str, text_b: str) -> None:
    sent_a = sentences(text_a)
    sent_b = {n for n, _ in sentences(text_b)}
    shared = [(n, original) for n, original in sent_a if n in sent_b]
    sh_a, sh_b = shingles(text_a), shingles(text_b)
    print(f"{name_a} → {name_b}")
    print(f"   содержательная часть: {len(normalize(text_a).split())} слов, {len(sent_a)} предложений")
    print(f"   дословно совпало: {len(shared)} ({100 * len(shared) / max(1, len(sent_a)):.1f} %)")
    print(f"   по 8-словным фрагментам: {100 * len(sh_a & sh_b) / max(1, len(sh_a)):.2f} %")
    for n, original in sorted(shared, key=lambda x: -len(x[0].split()))[:6]:
        print(f"      ({len(n.split())} сл.) {original[:130]}")
    print()


def main() -> None:
    texts = {name: body_only(document_text(path)) for name, path in DOCS.items()}
    names = list(DOCS)
    for i, name_a in enumerate(names):
        for name_b in names[i + 1:]:
            compare(name_a, texts[name_a], name_b, texts[name_b])


if __name__ == "__main__":
    main()
